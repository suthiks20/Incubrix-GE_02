"""Compare the generated main calendar to workbook tab 02_ASSETS."""
import csv
import sys
from pathlib import Path

from openpyxl import load_workbook

from incubrix_engine.export import COLS

ENGINE_ROOT = Path(__file__).resolve().parent
PROJECT_ROOT = ENGINE_ROOT.parent
WORKBOOK_PATH = PROJECT_ROOT / "IncuBrix_Option2_Submission_Workbook_FINAL_v2.xlsx"
CALENDAR_PATH = PROJECT_ROOT / "main_output" / "calendar.csv"
WORKSHEET_NAME = "02_ASSETS"
FIRST_WORKSHEET_ROW = 5


def cell_text(value):
    return "" if value is None else str(value)


def compare_rows(workbook_rows, csv_rows):
    mismatches = []
    for row_offset, (workbook_row, csv_row) in enumerate(zip(workbook_rows, csv_rows)):
        worksheet_row = FIRST_WORKSHEET_ROW + row_offset
        for column_index, column_name in enumerate(COLS):
            workbook_value = cell_text(workbook_row[column_index])
            csv_value = csv_row.get(column_name, "")
            if workbook_value != csv_value:
                mismatches.append(
                    f"row {worksheet_row}, column {column_name}: "
                    f"workbook={workbook_value!r}, calendar={csv_value!r}"
                )
    if len(workbook_rows) != len(csv_rows):
        mismatches.append(
            f"row count mismatch: workbook={len(workbook_rows)}, calendar={len(csv_rows)}"
        )
    return mismatches


def main():
    if not WORKBOOK_PATH.is_file():
        print(f"Workbook not found: {WORKBOOK_PATH}", file=sys.stderr)
        return 1
    if not CALENDAR_PATH.is_file():
        print(f"Calendar not found: {CALENDAR_PATH}", file=sys.stderr)
        return 1

    workbook = load_workbook(WORKBOOK_PATH, read_only=True, data_only=True)
    try:
        if WORKSHEET_NAME not in workbook.sheetnames:
            print(f"Worksheet not found: {WORKSHEET_NAME}", file=sys.stderr)
            return 1
        worksheet = workbook[WORKSHEET_NAME]
        workbook_rows = list(worksheet.iter_rows(
            min_row=FIRST_WORKSHEET_ROW,
            max_row=FIRST_WORKSHEET_ROW + 29,
            min_col=1,
            max_col=len(COLS),
            values_only=True,
        ))
    finally:
        workbook.close()

    with CALENDAR_PATH.open(encoding="utf-8-sig", newline="") as csv_file:
        csv_rows = list(csv.DictReader(csv_file))
    mismatches = compare_rows(workbook_rows, csv_rows)
    if mismatches:
        for mismatch in mismatches:
            print(mismatch)
        return 1
    print("Workbook and calendar match: 30 rows, columns A-U.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
